"""QQ bot stack component."""
















import io
import os
import re
import threading
import time
from collections import defaultdict
from datetime import datetime
from typing import Any, Dict, List, Optional

from .logger import setup_logger

logger = setup_logger("word_exporter")


DEFAULT_IMAGE_MAX = 800
DEFAULT_JPEG_QUALITY = 78

SMALL_IMAGE_BYTES = 80 * 1024

_LOCAL_PATH_RE = re.compile(r"^/api/images/\d+/(.+)$")





_BAD_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


def _xml_safe(text: Any) -> str:
    """QQ bot stack component."""
    if text is None:
        return ""
    if not isinstance(text, str):
        text = str(text)
    return _BAD_XML.sub("", text)


class WordExporter:
    """QQ bot stack component."""

    def __init__(self, export_dir: str, data_dir: str,
                 image_max: int = DEFAULT_IMAGE_MAX,
                 jpeg_quality: int = DEFAULT_JPEG_QUALITY,
                 bot_qq: str = ""):
        self.export_dir = os.path.abspath(export_dir)
        self.data_dir = os.path.abspath(data_dir)
        self.image_max = int(image_max)
        self.jpeg_quality = int(jpeg_quality)
        
        self.bot_qq = str(bot_qq or "")
        
        self._lock = threading.Lock()
        self.progress: Dict[str, Any] = {
            "running": False, "total": 0, "done": 0, "failed": 0,
            "messages": 0, "images": 0, "missing_images": 0,
            "out_dir": "", "error": "", "started_at": "", "finished_at": "",
            "current": "",
        }

    
    def _set(self, **kw):
        with self._lock:
            self.progress.update(kw)

    def _bump(self, **kw):
        with self._lock:
            for k, v in kw.items():
                self.progress[k] = self.progress.get(k, 0) + v

    def get_progress(self) -> Dict[str, Any]:
        with self._lock:
            return dict(self.progress)

    def _image_path(self, group_id: Any, seg: Dict[str, Any]) -> Optional[str]:
        """QQ bot stack component."""




        rel = str(seg.get("local_path") or "")
        m = _LOCAL_PATH_RE.match(rel)
        if m:
            tail = m.group(1).replace("/", os.sep)
            for cand in (os.path.join(self.export_dir, f"群{group_id}", tail),
                         os.path.join(self.data_dir, str(group_id), tail)):
                if os.path.isfile(cand):
                    return cand
        
        fname = str(seg.get("file") or "")
        if fname:
            for root in (os.path.join(self.export_dir, f"群{group_id}", "附件"),
                         os.path.join(self.data_dir, str(group_id), "附件")):
                if not os.path.isdir(root):
                    continue
                for dirpath, _dirs, files in os.walk(root):
                    if fname in files:
                        return os.path.join(dirpath, fname)
        return None

    def _image_bytes(self, path: str) -> Optional[bytes]:
        """QQ bot stack component."""
        try:
            size = os.path.getsize(path)
        except OSError:
            return None
        raw = None
        if size <= SMALL_IMAGE_BYTES and path.lower().endswith((".jpg", ".jpeg", ".png")):
            try:
                with open(path, "rb") as f:
                    return f.read()          
            except OSError:
                return None
        try:
            from PIL import Image
            with Image.open(path) as im:
                im = im.convert("RGB") if im.mode in ("RGBA", "P", "LA") else im
                if max(im.size) > self.image_max:
                    im.thumbnail((self.image_max, self.image_max))
                buf = io.BytesIO()
                im.convert("RGB").save(buf, format="JPEG", quality=self.jpeg_quality,
                                       optimize=True)
                raw = buf.getvalue()
        except Exception as e:
            logger.debug(f"缩图失败，回退原图: {path} ({type(e).__name__}: {e})")
            try:
                with open(path, "rb") as f:
                    raw = f.read()
            except OSError:
                return None
        return raw

    @staticmethod
    def _message_text(msg: Dict[str, Any]) -> str:
        """QQ bot stack component."""
        c = msg.get("content") or {}
        txt = str(c.get("text") or "").strip()
        if txt:
            return _xml_safe(txt)
        parts = []
        for seg in c.get("segments") or []:
            parts.append(str(seg.get("summary") or seg.get("text") or ""))
        return _xml_safe("".join(parts).strip())

    @staticmethod
    def _setup_doc():
        """QQ bot stack component."""
        from docx import Document
        from docx.shared import Pt
        from docx.oxml.ns import qn

        doc = Document()
        try:
            style = doc.styles["Normal"]
            style.font.name = "微软雅黑"
            style.font.size = Pt(10.5)
            rpr = style.element.get_or_add_rPr()
            rpr.get_or_add_rFonts().set(qn("w:eastAsia"), "微软雅黑")
        except Exception as e:
            logger.debug(f"设置字体失败（不影响导出）: {e}")
        return doc

    
    def _build_day(self, group_id: Any, date_part: str, msgs: List[Dict[str, Any]]):
        """QQ bot stack component."""




        from docx.shared import Inches, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        n_msg = n_img = n_missing = 0
        doc = self._setup_doc()
        head = doc.add_paragraph()
        hr = head.add_run(f"群 {group_id} · {date_part}　（共 {len(msgs)} 条）")
        hr.bold = True
        hr.font.size = Pt(14)

        for msg in msgs:
            who = _xml_safe(msg.get("card") or msg.get("nickname") or msg.get("user_id") or "")
            
            
            if msg.get("is_bot") or (self.bot_qq and str(msg.get("user_id")) == self.bot_qq):
                who = f"{who}[AI]"
            uid = _xml_safe(msg.get("user_id", ""))
            dt = str(msg.get("datetime") or "")
            clock = dt[11:19] if len(dt) >= 19 else dt
            text = self._message_text(msg)

            p = doc.add_paragraph()
            r1 = p.add_run(f"[{clock}] {who}(QQ:{uid}): ")
            r1.bold = True
            p.add_run(text)
            n_msg += 1

            
            for seg in ((msg.get("content") or {}).get("segments") or []):
                if seg.get("type") != "image":
                    continue
                src = self._image_path(group_id, seg)
                data = self._image_bytes(src) if src else None
                if not data:
                    n_missing += 1
                    doc.add_paragraph().add_run("[图片：本地没有副本，没能嵌进来]")
                    continue
                try:
                    pic = doc.add_paragraph()
                    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    pic.add_run().add_picture(io.BytesIO(data), width=Inches(2.4))
                    n_img += 1
                except Exception as e:
                    logger.debug(f"嵌图失败: {type(e).__name__}: {e}")
                    n_missing += 1
                    doc.add_paragraph().add_run("[图片：格式不支持，没能嵌进来]")
        return doc, n_msg, n_img, n_missing

    
    def export_all(self, store, groups: Optional[List[Any]] = None) -> Dict[str, Any]:
        """QQ bot stack component."""



        out_root = os.path.join(self.export_dir, "Word")
        self._set(running=True, total=0, done=0, failed=0, messages=0, images=0,
                  missing_images=0, error="", out_dir=out_root, current="",
                  started_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                  finished_at="")

        try:
            
            
            
            by_day: Dict[Any, List[Dict[str, Any]]] = defaultdict(list)
            seen = set()
            for group_id, msg in store.iter_all_messages():
                if groups and group_id not in groups:
                    continue
                if msg.get("type") != "message":
                    continue
                mid = msg.get("message_id", 0)
                key = ("id", group_id, mid) if mid else (
                    "fallback", group_id, msg.get("user_id", 0), msg.get("time", 0),
                    (msg.get("content") or {}).get("text", ""))
                if key in seen:
                    continue
                seen.add(key)
                dt = str(msg.get("datetime") or "")
                date_part = dt[:10] if len(dt) >= 10 else datetime.now().strftime("%Y-%m-%d")
                by_day[(group_id, date_part)].append(msg)

            self._set(total=len(by_day))
            logger.info(f"[Word导出] 共 {len(by_day)} 个「群+日期」要写")

            
            for (group_id, date_part), msgs in sorted(
                    by_day.items(), key=lambda kv: (str(kv[0][0]), kv[0][1])):
                
                
                
                msgs.sort(key=lambda m: (m.get("time", 0) or 0, str(m.get("message_id", ""))))
                self._set(current=f"群{group_id} {date_part}")
                out_dir = os.path.join(out_root, f"群{group_id}")
                out_path = os.path.join(out_dir, f"{date_part}.docx")
                try:
                    doc, n_msg, n_img, n_missing = self._build_day(group_id, date_part, msgs)
                    os.makedirs(out_dir, exist_ok=True)
                    doc.save(out_path)
                    self._bump(done=1, messages=n_msg, images=n_img,
                               missing_images=n_missing)
                    logger.info(f"[Word导出] 写好 {out_path}（{n_msg} 条 / {n_img} 图）")
                except Exception as e:
                    import traceback
                    self._bump(done=1, failed=1)
                    logger.error(f"[Word导出] 群{group_id} {date_part} 写失败，跳过继续: "
                                 f"{type(e).__name__}: {e}\n{traceback.format_exc()}")

            fin = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            self._set(running=False, finished_at=fin, current="")
            result = self.get_progress()
            logger.info(f"[Word导出] 全部完成：{result['done']} 个文档"
                        f"（失败 {result.get('failed', 0)}）、"
                        f"{result['messages']} 条消息、{result['images']} 张图"
                        f"（缺图 {result['missing_images']}）")
            return result
        except Exception as e:
            import traceback
            logger.error(f"[Word导出] 失败: {e}\n{traceback.format_exc()}")
            self._set(running=False, error=f"{type(e).__name__}: {e}",
                      finished_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
            return self.get_progress()
